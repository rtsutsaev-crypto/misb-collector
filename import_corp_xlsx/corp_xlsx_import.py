#!/usr/bin/env python3
"""Corporate leads workbook («Корпоративные лиды МИСБ — 20 отраслевых сегментов», срез 29.09.2026)
-> additions to the company directory and the source candidates.

The workbook is a re-cut of the data the directory already holds (368 groups, 1 338 INN, 933 sites,
151 channels come from the same monitor snapshot). What it adds and this script carries over:
  * the 20 industry segments (sheet «Отрасли» and column «Сегмент» of «Группы») with the MISB search themes
    of every segment  -> document orgdir/segments (also group site, procurement page, tier, supplier portal,
    corporate university / academy found in the registry and in lead customers);
  * INN hints and segments of the group candidates (sheet «Кандидаты 5000») -> orgdir/group-candidates
    (fields innHint and segment; `inn` stays empty so the EGRUL stage still resolves them);
  * procurement pages not yet known -> list for the verify queue;
  * tender rows of «Лиды МИСБ» that are not in the base yet -> leadset.

Usage:
  pip install openpyxl
  python3 corp_xlsx_import.py FILE.xlsx W [--date YYYY-MM-DD]

W: work dir with the database dump: W/db/orgdir/*.json (hier-N, orgs-N, misc, group-candidates),
W/db/leadsets/*.json, W/db/config/sources-plan.json. Writes W/out/orgdir-segments.json,
group-candidates-upd.json, xlsx-urls.json, xlsx-leads.json, xlsx-summary.json.
"""
from __future__ import annotations

import datetime as dt
import glob
import hashlib
import json
import os
import re
import sys
from urllib.parse import urlparse, parse_qs

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "b2b_center"))
from b2b_core import existing_leads, norm_title  # noqa: E402

CU_RE = re.compile(r"корпоративн\w*\s+(университет|академи|институт)", re.I)
CU_GROUP = {"росатом": "Росатом", "сбербанк": "Сбербанк", "норильск": "Норникель", "русгидро": "РусГидро", "газпром": "Газпром",
            "роснефт": "Роснефть", "лукойл": "ЛУКОЙЛ", "ржд": "РЖД"}
GROUP_ALIAS = {"Сбер": "Сбербанк"}  # the directory of organizations names the group differently from the workbook
TIER = {"A": 0, "B": 1, "C": 2}


def norm(s) -> str:
    return re.sub(r"[^а-яa-z0-9]", "", str(s or "").lower().replace("ё", "е"))


def nurl(u) -> str:
    return re.sub(r"^https?://(www\.)?", "", str(u or "").strip().lower()).rstrip("/")


def sheet_rows(wb, name):
    ws = wb[name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr = rows[0]
    return [dict(zip(hdr, r)) for r in rows[1:] if any(c is not None for c in r)]


def load_dir(W, prefix):
    out = []
    for p in sorted(glob.glob(os.path.join(W, "db", "orgdir", prefix + "-*.json"))):
        d = json.load(open(p, encoding="utf-8"))
        out += (d.get("data", d) or {}).get("rows", []) or []
    return out


def lead_id(url: str, title: str) -> str:
    q = parse_qs(urlparse(url).query)
    for k in ("regNumber", "regnumber"):
        if q.get(k):
            return q[k][0]
    m = re.search(r"tender-(\d+)", url)
    if m and "b2b-center" in url:
        return m.group(1)
    return "xl-" + hashlib.sha1((url + title).encode()).hexdigest()[:10]


def main(xlsx: str, W: str, today: str) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    G, Y = sheet_rows(wb, "Группы"), sheet_rows(wb, "Юрлица")
    K, L = sheet_rows(wb, "Кандидаты 5000"), sheet_rows(wb, "Лиды МИСБ")
    seg_rows = [r for r in [list(x) for x in wb["Отрасли"].iter_rows(values_only=True)]]
    hier, orgs = load_dir(W, "hier"), load_dir(W, "orgs")
    misc = json.load(open(os.path.join(W, "db", "orgdir", "misc.json"), encoding="utf-8"))
    misc = misc.get("data", misc)

    # --- segments: the summary table starts under the header «Отраслевой сегмент»
    segments, start = [], next(i for i, r in enumerate(seg_rows) if r and r[0] == "Отраслевой сегмент")
    for r in seg_rows[start + 1:]:
        if not r[0] or not r[1] or not str(r[1]).isdigit():
            continue
        segments.append({"name": r[0], "groups": int(r[1]), "goal": int(r[2] or 0), "missing": int(r[3] or 0), "withInn": int(r[4] or 0),
                         "withChannel": int(r[5] or 0), "children": int(r[6] or 0), "themes": r[7] or "", "candidates": int(r[8] or 0), "sheet": r[9] or r[0]})
    seg_name = {}  # «Сегмент» of the sheet «Группы» uses the same names as the summary table
    # --- best tier / group
    tier_of, orgs_of = {}, {}
    for o in orgs:
        g = o.get("g")
        if g:
            tier_of[g] = min(tier_of.get(g, 9), TIER.get(o.get("t"), 9))
            orgs_of.setdefault(g, []).append(o)
    portal = {}
    for p in misc.get("supplier_portals", []) or []:
        portal.setdefault(p.get("group_name"), {"url": p.get("url"), "title": p.get("title"), "mode": p.get("access_mode")})
    # corporate universities / academies: registry names and lead customers
    cu = {}
    for h in hier:
        if CU_RE.search(h.get("n", "")):
            cu.setdefault(h["g"], set()).add(h["n"])
    for o in orgs:
        if o.get("g") and CU_RE.search(o.get("n", "")):
            cu.setdefault(GROUP_ALIAS.get(o["g"], o["g"]), set()).add(o["n"])
    cust = {}
    for p in glob.glob(os.path.join(W, "db", "leadsets", "*.json")):
        d = json.load(open(p, encoding="utf-8"))
        for l in (d.get("data", d) or {}).get("leads", []) or []:
            c = str(l.get("customer") or "")
            if CU_RE.search(c):
                cust[re.sub(r"\s+", " ", c).strip().strip('"«»')] = cust.get(re.sub(r"\s+", " ", c).strip().strip('"«»'), 0) + 1
    for name, n in cust.items():
        for key, g in CU_GROUP.items():
            if key in name.lower():
                cu.setdefault(g, set()).add(name)
    by_inn, by_group = {}, {}
    for g in G:
        name = g["Группа"]
        row = {"g": name, "s": g["Сегмент"], "n": g["№ внутри сегмента*"], "site": g["Сайт"] or "", "proc": g["Закупочный URL"] or "",
               "ps": g["Статус URL"] or "", "ch": g["Дочек в файле"]}
        if name in tier_of:
            row["t"] = "ABC"[tier_of[name]]
        if name in portal:
            row["sp"] = portal[name]
        if name in cu:
            row["cu"] = sorted(cu[name])[:5]
        row = {k: v for k, v in row.items() if v not in ("", None)}
        by_inn[str(g["ИНН"])] = row
        by_group[norm(name)] = str(g["ИНН"])
    seg_doc = {"importedAt": today, "origin": "Корпоративные лиды МИСБ — 20 отраслевых сегментов, срез 29.09.2026 (xlsx пользователя)",
               "note": "Сегменты — укрупнённые коммерческие отрасли исходной подборки, не рейтинг по ВВП; № внутри сегмента — порядок отобранных строк, не независимое ранжирование по выручке. Ключ byInn — ИНН головного юрлица группы, byGroup — нормализованное имя группы. cu — корпоративные университеты и академии группы, найденные в справочнике и среди заказчиков лидов.",
               "segments": segments, "byInn": by_inn, "byGroup": by_group}

    # --- group candidates: INN hints and segments
    kn = {norm(k["Наименование"]): k for k in K}
    cand_path = os.path.join(W, "db", "orgdir", "group-candidates.json")
    cd = json.load(open(cand_path, encoding="utf-8")); cdata = cd.get("data", cd)
    n_inn = n_seg = 0
    for c in cdata["candidates"]:
        k = kn.get(norm(c["name"]))
        if not k:
            continue
        if k.get("ИНН по базе*") and not c.get("innHint"):
            c["innHint"] = str(k["ИНН по базе*"]); n_inn += 1
        if k.get("Отрасль группы*") and not c.get("segment"):
            c["segment"] = k["Отрасль группы*"]; n_seg += 1
        if k.get("Группа по базе") and not c.get("groupHint"):
            c["groupHint"] = k["Группа по базе"]
    cdata["note"] = (cdata.get("note", "") + f" 29.09.2026 из файла отраслей добавлены innHint (сверить по ЕГРЮЛ) и segment.").strip()

    # --- procurement pages not known yet
    known = set()
    for o in orgs:
        known.add(nurl(o.get("url")))
    for p in misc.get("priority_100", []) or []:
        known.add(nurl(p.get("procurement_url")))
    plan = json.load(open(os.path.join(W, "db", "config", "sources-plan.json"), encoding="utf-8"))
    for s in plan.get("sources", []):
        for u in s.get("urls", []) or []:
            known.add(nurl(u))
        known.add(nurl(s.get("from")))
    urls, seen = [], set()
    for src, key, ttl in [(G, "Закупочный URL", "группа"), (Y, "Собственный закупочный URL", "юрлицо")]:
        for r in src:
            u = r.get(key)
            if u and nurl(u) not in known and nurl(u) not in seen:
                seen.add(nurl(u))
                nm = r.get("Группа") if ttl == "группа" else f"{r.get('Юрлицо')} ({r.get('Группа')})"
                urls.append({"name": f"Закупки: {nm}", "section": f"закупочная страница ({ttl}), статус в исходном файле {r.get('Статус URL') or r.get('Статус канала') or '—'}", "url": u})

    # --- leads not in the base
    ids, pairs = existing_leads(W)
    base_urls = set()
    for p in glob.glob(os.path.join(W, "db", "leadsets", "*.json")):
        d = json.load(open(p, encoding="utf-8"))
        for l in (d.get("data", d) or {}).get("leads", []) or []:
            base_urls.add(nurl(l.get("url")))
    leads, skipped = [], {"уже в базе": 0, "без срока": 0}
    for r in L:
        url = str(r["Ссылка на закупку"] or "").strip()
        title = " ".join(str(r["Предмет"] or "").split())
        dl = str(r["Срок подачи"] or "")[:10]
        lid = lead_id(url, title)
        if nurl(url) in base_urls or lid in ids or (norm_title(title), dl) in pairs or (norm_title(title), "") in pairs:
            skipped["уже в базе"] += 1
            continue
        ids.add(lid)
        lead = {"id": lid, "title": title, "customer": str(r["Заказчик в карточке"] or "").strip(), "region": "" if r["Регион"] in (None, "Россия") else str(r["Регион"]),
                "price": float(r["Сумма"]) if r["Сумма"] not in (None, "") else None, "deadline": dl, "law": "", "url": url,
                "source": "mtrf:" + str(r["Источник"] or "прочее"), "collectedAt": today, "flags": [], "country": "RU", "currency": "RUB",
                "note": f"Файл отраслей 29.09.2026, лист «Лиды МИСБ»: {r['Категория'] or 'без категории'}, fit монитора {r['Fit монитора']}; {r['Актуальность на 29.09.2026'] or ''}"}
        if r["ИНН по базе*"]:
            lead["customerInn"] = str(r["ИНН по базе*"])
        leads.append(lead)

    out = os.path.join(W, "out")
    os.makedirs(out, exist_ok=True)
    json.dump(seg_doc, open(os.path.join(out, "orgdir-segments.json"), "w", encoding="utf-8"), ensure_ascii=False)
    json.dump(cdata, open(os.path.join(out, "group-candidates-upd.json"), "w", encoding="utf-8"), ensure_ascii=False)
    json.dump(urls, open(os.path.join(out, "xlsx-urls.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump({"source": "Файл отраслей 29.09.2026", "collectedAt": today, "leads": leads}, open(os.path.join(out, "xlsx-leads.json"), "w", encoding="utf-8"), ensure_ascii=False)
    summ = {"segments": len(segments), "groups": len(G), "entities": len(Y), "candidates_inn_hint": n_inn, "candidates_segment": n_seg,
            "new_urls": len(urls), "leads_new": len(leads), "leads_skipped": skipped, "groups_with_cu": len(cu)}
    json.dump(summ, open(os.path.join(out, "xlsx-summary.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return summ


if __name__ == "__main__":
    a = sys.argv[1:]
    if len(a) < 2:
        sys.exit(__doc__)
    day = a[a.index("--date") + 1] if "--date" in a else dt.date.today().isoformat()
    print(json.dumps(main(a[0], a[1], day), ensure_ascii=False, indent=1))
